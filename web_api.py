# -*- coding: utf-8 -*-
"""
Мост между страницей (index.html) и pokazaniya.py. Работает внутри Pyodide (Python в браузере),
но проверяется и обычным Python: рабочая папка берётся из POKAZ_WORK (в браузере это /work).

Страница кладёт файлы в <WORK>/in (и необязательно в <WORK>/add), вызывает process(режим, дата)
и получает JSON с итогами; готовый zip лежит в <WORK>/_result.zip.
"""
import contextlib
import io
import json
import os
import shutil
import sys
import tempfile
import traceback
import zipfile

import pokazaniya as P

WORK = os.environ.get("POKAZ_WORK", "/work")
IN = os.path.join(WORK, "in")
ADD = os.path.join(WORK, "add")
RESULT = os.path.join(WORK, "_result.zip")
LIMIT = 300          # сколько строк замечаний отдавать странице


def reset():
    """Очистить рабочую папку перед новым запуском."""
    shutil.rmtree(WORK, ignore_errors=True)
    os.makedirs(IN)
    os.makedirs(ADD)
    os.chdir(WORK)


def _run(argv):
    """Запустить pokazaniya.py как из командной строки. -> (код, вывод)."""
    buf = io.StringIO()
    old = sys.argv
    sys.argv = ["pokazaniya.py"] + list(argv)
    code = 0
    try:
        with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(buf):
            P.main()
    except SystemExit as e:
        if isinstance(e.code, str):
            buf.write(e.code + "\n")
            code = 1
        elif e.code not in (None, 0):
            code = e.code
    except Exception:
        buf.write(traceback.format_exc())
        code = 1
    finally:
        sys.argv = old
    return code, buf.getvalue()


def _clean_log(text):
    """Убрать из журнала то, что имеет смысл только в командной строке (пути, ключи)."""
    out = []
    for line in text.splitlines():
        if line.startswith(("Отчёт:", "Сводная таблица:", "Сводка:")):
            continue
        line = line.replace("Исправьте показания в файлах прошлого месяца и запустите команду снова с --overwrite.",
                            "Исправьте показания в таблицах прошлого месяца и загрузите их заново.")
        line = line.replace("уберите «?» в таблице прошлого месяца и запустите снова с --overwrite.",
                            "уберите «?» в таблице прошлого месяца и загрузите её заново.")
        if line.startswith(("Новый месяц", "Готово:")) and ". Папка:" in line:
            line = line.split(". Папка:")[0] + "."
        out.append(line)
    return "\n".join(out).strip()


def _issue_dict(i):
    r = i.row
    prev = P.num(r.prev)
    return {
        "level": i.level, "kind": i.kind, "podezd": P.slug(r.title), "res": r.res,
        "flat": P.short(r.address), "text": i.text,
        "value": None if r.new is None else str(r.new),
        "prev": prev, "prev_date": P.fmt_date(r.prev_date),
        "file": os.path.basename(r.src),
    }


def _summary_rows(data):
    tmp = tempfile.mkdtemp()
    try:
        path = P.run_summary(data, tmp)
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True)["Подъезды"]
        rows = [[("" if v is None else v) for v in r] for r in ws.iter_rows(values_only=True)]
        return {"headers": rows[0], "rows": rows[1:]}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _zip(parts, name):
    """parts: [(папка на диске, префикс в архиве)]. Пишет RESULT, возвращает число файлов."""
    n = 0
    with zipfile.ZipFile(RESULT, "w", zipfile.ZIP_DEFLATED) as z:
        for folder, prefix in parts:
            for root, _, files in os.walk(folder):
                for f in sorted(files):
                    full = os.path.join(root, f)
                    rel = os.path.relpath(full, folder)
                    z.write(full, os.path.join(prefix, rel) if prefix else rel)
                    n += 1
    return n


def _new_month_dirs():
    return sorted(d for d in os.listdir(WORK)
                  if d.startswith("подъезды_") and os.path.isdir(os.path.join(WORK, d)))


def process(mode, date_str=""):
    """mode: close | check | split. date_str: ДД.ММ.ГГГГ или пусто. -> JSON-строка."""
    if os.path.exists(RESULT):
        os.remove(RESULT)
    inputs = [f for f in os.listdir(IN) if not f.startswith(("~$", "_"))]
    if not inputs:
        return json.dumps({"ok": False, "log": "Не загружено ни одного файла."}, ensure_ascii=False)
    has_add = bool(os.listdir(ADD))
    if mode == "close":
        argv = ["close", IN] + (["--date", date_str] if date_str else []) + (["--add", ADD] if has_add else [])
    elif mode == "check":
        argv = ["check", IN]
    elif mode == "split":
        argv = ["split", IN] + (["--date", date_str] if date_str else [])
    else:
        return json.dumps({"ok": False, "log": "Неизвестный режим: %s" % mode}, ensure_ascii=False)

    code, log = _run(argv)
    res = {"ok": code == 0, "mode": mode, "log": _clean_log(log), "zip_name": None, "files": 0}
    if code != 0:
        return json.dumps(res, ensure_ascii=False)

    parts, zip_name = [], None
    reports = os.path.join(IN, P.REPORT_DIR)
    months = _new_month_dirs()
    if mode == "close":
        parts = [(os.path.join(WORK, d), d) for d in months] + [(reports, P.REPORT_DIR)]
        zip_name = "результат_%s.zip" % (months[0] if months else "месяц")
    elif mode == "check":
        parts = [(reports, P.REPORT_DIR)]
        zip_name = "проверка.zip"
    else:
        parts = [(os.path.join(WORK, d), d) for d in months]
        zip_name = "%s.zip" % (months[0] if months else "подъезды")
    res["files"] = _zip([p for p in parts if os.path.isdir(p[0])], zip_name)
    res["zip_name"] = zip_name

    if mode in ("close", "check"):
        with contextlib.redirect_stdout(io.StringIO()):
            data, nfiles = P.load_inputs([IN])
            issues = P.run_check(data)
        n_err, n_warn, n_doubt, n_nomark = P.count_kinds(issues)
        res["counts"] = {"errors": n_err, "warnings": n_warn, "doubts": n_doubt, "nomarks": n_nomark,
                         "other": n_warn - n_doubt,
                         "tables": nfiles, "rows": sum(1 for _ in data.rows())}
        res["errors"] = [_issue_dict(i) for i in issues if i.level == "ОШИБКА"][:LIMIT]
        res["doubts"] = [_issue_dict(i) for i in issues if i.kind == "doubt"][:LIMIT]
        res["warnings"] = [_issue_dict(i) for i in issues
                           if i.level != "ОШИБКА" and i.kind != "doubt"][:LIMIT]
        res["summary"] = _summary_rows(data)
    return json.dumps(res, ensure_ascii=False)


def selftest():
    """Мини-прогон на встроенной таблице: проверяет, что модуль в браузере работает целиком."""
    from datetime import datetime
    from openpyxl import Workbook
    reset()
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "ГВС"
        ws["A1"] = "дом 1 подъезд 1"
        rows = [("ул.Тестовая, д.1, кв.1", 105, 100, "A1"),
                ("ул.Тестовая, д.1, кв.2", "50?", 40, "A2"),
                ("ул.Тестовая, д.1, кв.3", 90, 100, "A3")]
        for i, (a, new, prev, serial) in enumerate(rows, start=2):
            for j, v in enumerate([a, "ГВС", new, prev, serial, datetime(2026, 10, 19), datetime(2026, 9, 21)], 1):
                ws.cell(i, j, v)
        wb.save(os.path.join(IN, "тест.xlsx"))
        res = json.loads(process("close", "20.11.2026"))
        with zipfile.ZipFile(RESULT) as z:
            names = z.namelist()
        ok = (res.get("ok") and res["counts"]["errors"] == 1 and res["counts"]["doubts"] == 1
              and "подъезды_2026-11/д1_п1.xlsx" in names)
        detail = "" if ok else "неожиданный результат: %s" % json.dumps(res.get("counts"), ensure_ascii=False)
    except Exception:
        ok, detail = False, traceback.format_exc()
    finally:
        reset()
    return json.dumps({"ok": bool(ok), "detail": detail}, ensure_ascii=False)
