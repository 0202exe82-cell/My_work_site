#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Показания приборов учёта (ГВС, ХВС, ОТ).

  split     выгрузки (.xls/.xlsx) -> файлы по подъездам
  check     проверка заполненных файлов по подъездам
  summary   сводка расхода по подъездам / корпусам / домам
  assemble  обратная сборка в общие таблицы для бухгалтерии
  next      файлы на следующий месяц из таблиц прошлого месяца + дата
  close     всё сразу: сводные, сводка, следующий месяц; ошибки — в конце

Подробности: python pokazaniya.py --help   и файл ИНСТРУКЦИЯ.txt
"""
import argparse
import math
import os
import re
import statistics
import sys
from collections import Counter, OrderedDict, defaultdict
from dataclasses import dataclass
from datetime import date, datetime

try:
    sys.stdout.reconfigure(errors="replace")
except Exception:
    pass

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Не установлена библиотека openpyxl. Выполните: pip install openpyxl xlrd")

# ----------------------------------------------------------------------------
# НАСТРОЙКИ (можно менять)
# ----------------------------------------------------------------------------
RESOURCES = ["ГВС", "ХВС", "ОТ"]          # порядок листов
JUMP_FACTOR = 5                            # «скачок»: расход больше типичного в N раз ...
JUMP_MIN = {"ГВС": 8, "ХВС": 12, "ОТ": 3}  # ... и больше этого значения (в единицах прибора)
ZERO_WARN = {"ГВС": False, "ХВС": False, "ОТ": False}  # нулевой расход — норма (True = предупреждать)
FULL_DECIMALS = 3                          # знаков в «полном» показании. Если предыдущее показание в выгрузке
                                           # записано короче (целое, 15.2, 0.65 — округлено), то новое
                                           # показание чуть ниже него (в пределах последнего знака) — не ошибка
DOUBT_CHARS = ("?", "？")                   # знак вопроса в новом показании = «сомнительно», не принимается
NO_READING_MARKS = ("н",)                  # отметка «показание не снято» (как пустое, но с пояснением)
MIN_GROUP = 5                              # минимум квартир в подъезде для своей «типичной» величины
NEXT_DAY = 20                              # дата съёма нового месяца: это число следующего месяца
REPORT_DIR = "_итоги"                      # подпапка для отчётов внутри папки месяца

# ----------------------------------------------------------------------------
# Оформление
# ----------------------------------------------------------------------------
F_BODY = Font(name="Trebuchet MS", size=9)
F_HEAD = Font(name="Trebuchet MS", size=10, bold=True)
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM = "#0.0000"
DATEF = "dd\\.mm\\.yyyy"
FILL_ERR = PatternFill("solid", fgColor="F8CBAD")
FILL_WARN = PatternFill("solid", fgColor="FFF2CC")


# ----------------------------------------------------------------------------
# Модель данных
# ----------------------------------------------------------------------------
@dataclass
class Row:
    title: str          # «дом 5 корпус 4а подъезд 3»
    res: str            # ГВС / ХВС / ОТ
    address: str
    serial: object
    prev_date: object
    prev: object        # предыдущее показание
    new_date: object    # дата съёма
    new: object         # новое показание
    src: str = ""
    line: int = 0


class Data:
    def __init__(self):
        self.entr = OrderedDict()   # title -> OrderedDict(res -> [Row])

    def add(self, r):
        self.entr.setdefault(r.title, OrderedDict()).setdefault(r.res, []).append(r)

    def titles(self):
        return sorted(self.entr, key=title_key)

    def rows(self):
        for t in self.titles():
            for res in res_order(self.entr[t]):
                for r in self.entr[t][res]:
                    yield r


def res_order(keys):
    keys = list(keys)
    return [r for r in RESOURCES if r in keys] + [r for r in keys if r not in RESOURCES]


TITLE_RE = re.compile(r"дом\s+(\S+)(?:\s+корпус\s+(\S+))?(?:\s+подъезд\s+(\S+))?", re.I)


def parse_title(t):
    m = TITLE_RE.search(t or "")
    if not m:
        return None, "", ""
    return m.group(1), m.group(2) or "", m.group(3) or ""


def nat(s):
    return [(0, int(p)) if p.isdigit() else (1, p) for p in re.split(r"(\d+)", str(s).lower()) if p]


def title_key(t):
    h, c, e = parse_title(t)
    if h is None:
        return ([(2, t.lower())], [], [])
    return (nat(h), nat(c), nat(e))


def slug(title):
    h, c, e = parse_title(title)
    if h is None:
        return re.sub(r"[^\w\-]+", "_", title).strip("_") or "подъезд"
    s = "д" + h
    if c:
        s += "_к" + c
    if e:
        s += "_п" + e
    return s


# ----------------------------------------------------------------------------
# Вспомогательные
# ----------------------------------------------------------------------------
def to_dt(v):
    if isinstance(v, datetime):
        return v.replace(hour=0, minute=0, second=0, microsecond=0)
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, str):
        for f in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
            try:
                return datetime.strptime(v.strip(), f)
            except ValueError:
                pass
    return None


def date_arg(s):
    d = to_dt(s)
    if d is None:
        raise argparse.ArgumentTypeError("дата нужна в виде ДД.ММ.ГГГГ, например 19.11.2026")
    return d


def is_empty(v):
    return v is None or (isinstance(v, str) and not v.strip())


def num(v):
    """Число или None (понимает «12,5»)."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        x = float(v)
        return x if math.isfinite(x) else None
    if isinstance(v, str):
        s = v.strip().replace(" ", "").replace(",", ".")
        try:
            x = float(s)
            return x if math.isfinite(x) else None
        except ValueError:
            return None
    return None


def ser(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return s or None


def drop(n, p):
    """На сколько новое показание меньше предыдущего (0, если не меньше)."""
    return max(0.0, round(p - n, 6))


def decimals(v):
    t = ("%.10f" % float(v)).rstrip("0")
    return len(t.split(".")[1]) if "." in t else 0


def reading_ok(n, p):
    """Новое показание не меньше предыдущего (с допуском на округление предыдущего в выгрузке)."""
    d = drop(n, p)
    if d == 0:
        return True
    k = decimals(p)
    return k < FULL_DECIMALS and d < 10 ** (-k)


def is_doubtful(v):
    """В тексте показания есть знак вопроса, например «107?» — значение сомнительное."""
    return isinstance(v, str) and any(c in v for c in DOUBT_CHARS)


def is_no_reading(v):
    return isinstance(v, str) and v.strip().casefold() in NO_READING_MARKS


def accepted_reading(r):
    """Число, если новое показание можно принять как достоверное; иначе None.
    Не принимаются: пустое, со знаком «?», нечисловое, отрицательное, меньше предыдущего."""
    if is_doubtful(r.new):
        return None
    n, p = num(r.new), num(r.prev)
    if n is None or n < 0:
        return None
    if p is not None and (p < 0 or not reading_ok(n, p)):
        return None
    return n


def cons_of(r):
    n, p = accepted_reading(r), num(r.prev)
    if n is None or p is None:
        return None
    return max(0.0, n - p)


def short(addr):
    return str(addr).split(",")[-1].strip()


def collect_sources(paths):
    files = []
    for p in paths:
        if os.path.isdir(p):
            for f in sorted(os.listdir(p)):
                if f.lower().endswith((".xls", ".xlsx")) and not f.startswith(("~$", "_")):
                    files.append(os.path.join(p, f))
        elif os.path.isfile(p):
            files.append(p)
        else:
            sys.exit("Не найдено: %s" % p)
    if not files:
        sys.exit("Не нашёл ни одного файла .xls/.xlsx.")
    return files


# ----------------------------------------------------------------------------
# Чтение
# ----------------------------------------------------------------------------
def norm(v):
    return (v.strip() or None) if isinstance(v, str) else v


def load_sheets(path):
    """[(имя листа, строки)] — все листы файла .xlsx или .xls (даты -> datetime)."""
    low = path.lower()
    if low.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(path, data_only=True)
        return [(ws.title, [[norm(v) for v in r] for r in ws.iter_rows(values_only=True)])
                for ws in wb.worksheets]
    if low.endswith(".xls"):
        try:
            import xlrd
        except ImportError:
            sys.exit("Для чтения .xls нужна библиотека xlrd: pip install xlrd\n"
                     "(или откройте файл в Excel и сохраните как .xlsx)")
        book = xlrd.open_workbook(path)
        out = []
        for si in range(book.nsheets):
            sh = book.sheet_by_index(si)
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    t, v = sh.cell_type(r, c), sh.cell_value(r, c)
                    if t == xlrd.XL_CELL_DATE:
                        v = xlrd.xldate_as_datetime(v, book.datemode)
                    elif t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        v = None
                    elif t == xlrd.XL_CELL_TEXT:
                        v = v.strip() or None
                    row.append(v)
                rows.append(row)
            out.append((sh.name, rows))
        return out
    sys.exit("Неизвестный тип файла: %s" % path)


def load_rows(path):
    """Строки первого листа (для выгрузок)."""
    return load_sheets(path)[0][1]


def cell(row, i):
    return row[i] if i < len(row) else None


def parse_export(path):
    """Выгрузка: строка-заголовок подъезда, затем строки счётчиков.
    Колонки: адрес, тип, серийный №, дата пред., показание пред., дата съёма, (новое)."""
    out, title = [], None
    for i, row in enumerate(load_rows(path), 1):
        a, b = cell(row, 0), cell(row, 1)
        if a is None and b is None:
            continue
        if b is None:
            title = str(a).strip()
            continue
        if title is None:
            continue
        res = str(b).strip().upper()
        out.append(Row(title=title, res=res, address=str(a).strip(), serial=ser(cell(row, 2)),
                       prev_date=to_dt(cell(row, 3)), prev=cell(row, 4),
                       new_date=to_dt(cell(row, 5)), new=cell(row, 6),
                       src=path, line=i))
    return out


def serial_like(v):
    """Похоже ли значение на серийный номер счётчика (а не на показание)."""
    if isinstance(v, str):
        t = v.strip()
        return bool(re.search(r"[A-Za-zА-Яа-я]", t)) or len(re.sub(r"\D", "", t)) >= 6
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v).is_integer() and abs(v) >= 1_000_000
    return False


def serial_share(vals):
    vals = [v for v in vals if not is_empty(v)][:40]
    return sum(1 for v in vals if serial_like(v)) / len(vals) if vals else 0.0


def detect_layout(body):
    """Раскладка столбцов листа по содержимому:
    standard — C новое, D предыдущее, E серийный №, F дата съёма, G дата предыдущего;
    swapped  — то же, но D серийный №, E предыдущее показание;
    raw      — как в исходной выгрузке: C серийный №, D дата пред., E пред. показание, F дата, G новое."""
    c_share = serial_share([cell(r, 2) for r in body])
    d_vals = [cell(r, 3) for r in body]
    d_share = serial_share(d_vals)
    e_share = serial_share([cell(r, 4) for r in body])
    nonempty = [v for v in d_vals if not is_empty(v)]
    d_date = sum(1 for v in nonempty if to_dt(v) is not None) / len(nonempty) if nonempty else 0.0
    if c_share >= 0.6 and d_date >= 0.6:
        return "raw"
    if d_share >= 0.6 and e_share < 0.4:
        return "swapped"
    return "standard"


LAYOUT_NOTE = {
    "swapped": "серийный № в столбце D, предыдущее показание в E",
    "raw": "раскладка исходной выгрузки (серийный № в C, новое показание в G)",
}


def load_entrance_file(path, data, quiet=False):
    """Файл подъезда (.xlsx/.xls): листы ГВС/ХВС/ОТ; в A1 название подъезда.
    Раскладка столбцов определяется отдельно для каждого листа."""
    base = os.path.basename(path)
    for name, rows in load_sheets(path):
        res = name.strip().upper()
        if res not in RESOURCES:
            if not quiet:
                print("  (лист «%s» в файле %s пропущен — это не ГВС/ХВС/ОТ)" % (name, base))
            continue
        title = str(cell(rows[0], 0) or "").strip() if rows else ""
        if not title:
            sys.exit("В файле %s на листе %s пустая ячейка A1 (должно быть название подъезда)." % (base, name))
        body = rows[1:]
        layout = detect_layout(body)
        if layout != "standard" and not quiet:
            print("  Файл %s, лист %s: %s — прочитал так." % (base, name, LAYOUT_NOTE[layout]))
        for i, r in enumerate(body, start=2):
            a = cell(r, 0)
            if is_empty(a):
                continue
            if layout == "raw":
                new, prev, serial = cell(r, 6), cell(r, 4), cell(r, 2)
                new_date, prev_date = cell(r, 5), cell(r, 3)
            else:
                new, new_date, prev_date = cell(r, 2), cell(r, 5), cell(r, 6)
                prev, serial = (cell(r, 4), cell(r, 3)) if layout == "swapped" else (cell(r, 3), cell(r, 4))
            data.add(Row(title=title, res=res, address=str(a).strip(), new=new, prev=prev,
                         serial=ser(serial), new_date=to_dt(new_date), prev_date=to_dt(prev_date),
                         src=path, line=i))


def load_inputs(paths):
    """Таблицы подъездов: папка или отдельные файлы .xlsx/.xls (в каждом — разделы ГВС/ХВС/ОТ листами)."""
    files = []
    for p in paths:
        if os.path.isdir(p):
            files += [os.path.join(p, f) for f in sorted(os.listdir(p))
                      if f.lower().endswith((".xlsx", ".xls")) and not f.startswith(("~$", "_"))]
        elif os.path.isfile(p):
            files.append(p)
        else:
            sys.exit("Не найдено: %s" % p)
    if not files:
        sys.exit("Не нашёл ни одного файла .xlsx/.xls с таблицами подъездов.")
    data = Data()
    for f in files:
        load_entrance_file(f, data)

    # Один и тот же подъезд/ресурс/адрес из двух файлов обычно означает,
    # что в папку попали копии или таблицы разных месяцев. Не считаем их дважды.
    source_by_key = {}
    duplicates = []
    for r in data.rows():
        key = (r.title.strip().casefold(), r.res.strip().upper(), re.sub(r"\s+", " ", r.address).strip().casefold())
        src = os.path.realpath(r.src)
        previous_src = source_by_key.get(key)
        if previous_src is not None and previous_src != src:
            duplicates.append((r, previous_src))
        else:
            source_by_key[key] = src
    if duplicates:
        examples = []
        for r, old_src in duplicates[:8]:
            examples.append("  %s / %s / %s: %s и %s" %
                            (r.title, r.res, r.address, os.path.basename(old_src), os.path.basename(r.src)))
        sys.exit("ОСТАНОВЛЕНО: одинаковые адреса одного ресурса найдены в разных входных файлах.\n"
                 "Вероятно, смешаны копии или месяцы; сводка иначе посчитала бы строки дважды.\n" +
                 "Примеры:\n" + "\n".join(examples) +
                 "\nОставьте в папке один актуальный файл на подъезд и запустите снова.")
    return data, len(files)


def base_dir(paths):
    p = os.path.abspath(paths[0])
    return p if os.path.isdir(p) else os.path.dirname(p)


def default_report_dir(paths):
    return os.path.join(base_dir(paths), REPORT_DIR)


# ----------------------------------------------------------------------------
# Запись
# ----------------------------------------------------------------------------
def style_cell(c, h, nf=None):
    c.font = F_BODY
    c.border = BORDER
    c.alignment = Alignment(horizontal=h, vertical="top")
    if nf:
        c.number_format = nf


def flen(v):
    if v is None:
        return 0
    if isinstance(v, (datetime, date)):
        return 10
    if isinstance(v, (int, float)):
        return len("%.4f" % v)
    return len(str(v))


def fit(rows_values, minimum=6.0):
    m = max([flen(v) for v in rows_values] or [0])
    return max(minimum, round(m * 0.95 + 1.6, 2))


def write_entrance(path, title, by_res):
    wb = Workbook()
    wb.remove(wb.active)
    spec = [("left", None), ("left", None), ("right", NUM), ("right", NUM),
            ("left", None), ("center", DATEF), ("center", DATEF)]
    for res in res_order(by_res):
        rows = by_res[res]
        ws = wb.create_sheet(res[:31])
        ws["A1"] = title
        ws["A1"].font = F_HEAD
        for i, r in enumerate(rows, start=2):
            vals = [r.address, r.res, r.new, r.prev, r.serial, to_dt(r.new_date), to_dt(r.prev_date)]
            for j, v in enumerate(vals):
                c = ws.cell(i, j + 1, v)
                style_cell(c, *spec[j])
        a_len = max([len(r.address) for r in rows] or [0])
        widths = {"A": max(30.3, round(a_len * 0.86, 1)), "B": 5.0, "C": 13.28,
                  "D": fit([r.prev for r in rows], 8.0), "E": fit([r.serial for r in rows], 8.0),
                  "F": 9.7, "G": 9.7}
        for k, w in widths.items():
            ws.column_dimensions[k].width = w
    wb.save(path)


def write_folder(data, out):
    n = 0
    for t in data.titles():
        write_entrance(os.path.join(out, slug(t) + ".xlsx"), t, data.entr[t])
        n += 1
    return n


def ensure_out(out, overwrite):
    if os.path.isdir(out) and not overwrite:
        old = [f for f in os.listdir(out) if f.lower().endswith(".xlsx") and not f.startswith(("~$", "_"))]
        if old:
            sys.exit("Папка %s уже существует и в ней есть файлы.\n"
                     "Укажите другую через --out или добавьте --overwrite." % out)
    os.makedirs(out, exist_ok=True)


def ensure_report_dir(outdir, overwrite, kinds):
    """Не заменять уже существующие отчёты без явного --overwrite."""
    if overwrite or not os.path.isdir(outdir):
        return
    old = []
    for f in sorted(os.listdir(outdir)):
        if (("check" in kinds and f == "проверка.xlsx") or ("summary" in kinds and f == "сводка.xlsx")
                or ("assemble" in kinds and f.endswith("_сводная.xlsx"))):
            old.append(f)
    if old:
        sys.exit("В папке %s уже есть отчёты: %s.\nЧтобы заменить их, добавьте --overwrite "
                 "или укажите другую папку через --out." % (outdir, ", ".join(old)))


def save_book(path, sheets):
    """sheets: [{name, headers, rows, widths, numfmt:{col_idx:fmt}, fill: fn(row)->Fill, bold_last}]"""
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(sh["name"][:31])
        ws.append(sh["headers"])
        for c in ws[1]:
            c.font = F_HEAD
        for row in sh["rows"]:
            ws.append(row)
        for i, w in enumerate(sh["widths"], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for rr in range(2, ws.max_row + 1):
            for ci, fmt in (sh.get("numfmt") or {}).items():
                ws.cell(rr, ci).number_format = fmt
            if sh.get("fill"):
                f = sh["fill"]([ws.cell(rr, c).value for c in range(1, len(sh["headers"]) + 1)])
                if f:
                    for c in range(1, len(sh["headers"]) + 1):
                        ws.cell(rr, c).fill = f
        if sh.get("bold_last") and ws.max_row > 1:
            for c in ws[ws.max_row]:
                c.font = F_HEAD
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
    wb.save(path)


# ----------------------------------------------------------------------------
# Проверка
# ----------------------------------------------------------------------------
@dataclass
class Issue:
    level: str
    row: Row
    text: str
    cons: object = None
    kind: str = ""      # "doubt" — значение с «?», "nomark" — отметка «н», иначе обычная ошибка/предупреждение


def run_check(data):
    issues, items = [], []
    seen = set()
    by_group, by_res = defaultdict(list), defaultdict(list)
    for r in data.rows():
        key = (r.title, r.res, r.address.lower())
        if key in seen:
            issues.append(Issue("ПРЕДУПРЕЖДЕНИЕ", r, "повтор: этот адрес уже есть на листе"))
        seen.add(key)
        if is_empty(r.new):
            issues.append(Issue("ПРЕДУПРЕЖДЕНИЕ", r, "показание не заполнено"))
            continue
        if is_doubtful(r.new):
            issues.append(Issue("ПРЕДУПРЕЖДЕНИЕ", r, "сомнительное показание со знаком «?»: «%s» — не принято, "
                                "перенесено предыдущее" % r.new, kind="doubt"))
            continue
        if is_no_reading(r.new):
            issues.append(Issue("ПРЕДУПРЕЖДЕНИЕ", r, "показание не снято (отметка «%s») — перенесено предыдущее"
                                % r.new, kind="nomark"))
            continue
        n, p = num(r.new), num(r.prev)
        if n is None:
            issues.append(Issue("ОШИБКА", r, "в ячейке не число или не конечное число: «%s»" % r.new))
            continue
        if n < 0:
            issues.append(Issue("ОШИБКА", r, "отрицательное показание: %g" % n))
            continue
        if p is not None and p < 0:
            issues.append(Issue("ОШИБКА", r, "отрицательное предыдущее показание: %g" % p))
            continue
        if p is None:
            issues.append(Issue("ПРЕДУПРЕЖДЕНИЕ", r, "нет предыдущего показания — расход не посчитать"))
            continue
        if not reading_ok(n, p):
            issues.append(Issue("ОШИБКА", r, "показание меньше предыдущего (%g < %g): опечатка, "
                                "замена счётчика или переход через ноль" % (n, p)))
            continue
        c = max(0.0, n - p)     # при округлённом предыдущем небольшой «минус» считаем нулевым расходом
        items.append((r, c))
        if c > 0:
            by_group[(r.title, r.res)].append(c)
            by_res[r.res].append(c)
        elif ZERO_WARN.get(r.res, True):
            issues.append(Issue("ПРЕДУПРЕЖДЕНИЕ", r, "нулевой расход", 0.0))
    med_res = {k: statistics.median(v) for k, v in by_res.items() if v}
    for r, c in items:
        grp = by_group.get((r.title, r.res), [])
        med = statistics.median(grp) if len(grp) >= MIN_GROUP else med_res.get(r.res)
        if med is None:
            continue
        if c > max(JUMP_FACTOR * med, JUMP_MIN.get(r.res, 0)):
            issues.append(Issue("ПРЕДУПРЕЖДЕНИЕ", r,
                                "расход %.4g заметно выше обычного (типично около %.4g)" % (c, med), c))
    for i in issues:
        if i.cons is None:
            i.cons = cons_of(i.row)
    issues.sort(key=lambda i: (i.level != "ОШИБКА", title_key(i.row.title), i.row.res, nat(i.row.address)))
    return issues


def fmt_date(d):
    d = to_dt(d)
    return d.strftime("%d.%m.%Y") if d else "—"


def print_doubtful(issues, limit=60):
    """Отдельное предупреждение о значениях со знаком «?» и о том, что перенесено вместо них."""
    doubt = [i for i in issues if i.kind == "doubt"]
    if not doubt:
        return
    print("\n" + "-" * 60)
    print("СОМНИТЕЛЬНЫЕ ПОКАЗАНИЯ («?»): %d. Они НЕ приняты; в новый месяц перенесено предыдущее показание"
          " и его дата." % len(doubt))
    for i in doubt[:limit]:
        r = i.row
        print("  %s | %s | %s | %s: «%s» -> перенесено %s (от %s)" %
              (os.path.basename(r.src), slug(r.title), r.res, short(r.address), r.new,
               "—" if is_empty(r.prev) else "%g" % num(r.prev) if num(r.prev) is not None else r.prev,
               fmt_date(r.prev_date)))
    if len(doubt) > limit:
        print("  ... и ещё %d (см. лист «Сомнительные» в отчёте)" % (len(doubt) - limit))
    print("Уточните показания, уберите «?» в таблице прошлого месяца и запустите снова с --overwrite.")


def count_kinds(issues):
    return (sum(1 for i in issues if i.level == "ОШИБКА"),
            sum(1 for i in issues if i.level != "ОШИБКА"),
            sum(1 for i in issues if i.kind == "doubt"),
            sum(1 for i in issues if i.kind == "nomark"))


def print_issues(issues, limit=25, final=False):
    errs = [i for i in issues if i.level == "ОШИБКА"]
    n_err, n_warn, n_doubt, n_nomark = count_kinds(issues)
    extra = ""
    if n_doubt or n_nomark:
        extra = ", из них со знаком «?»: %d, с отметкой «н»: %d" % (n_doubt, n_nomark)
    if final:
        print("\n" + "=" * 60)
        if not errs:
            print("Ошибок в данных нет. Предупреждений: %d%s." % (n_warn, extra))
        else:
            print("ВНИМАНИЕ: в данных найдено ошибок: %d (предупреждений: %d%s)." % (n_err, n_warn, extra))
            print("Эти строки перенесены в новый месяц БЕЗ нового показания (остались старые).")
    else:
        print("Проверка: ошибок %d, предупреждений %d%s." % (n_err, n_warn, extra))
    for i in errs[:limit]:
        print("  ОШИБКА  %s  %s  %s: %s" % (slug(i.row.title), i.row.res, short(i.row.address), i.text))
    if len(errs) > limit:
        print("  ... и ещё %d ошибок (см. отчёт)" % (len(errs) - limit))
    if final:
        if errs:
            print("Исправьте показания в файлах прошлого месяца и запустите команду снова с --overwrite.")
        print_doubtful(issues)
    elif any(i.kind == "doubt" for i in issues):
        print_doubtful(issues)


def write_check_report(issues, outdir):
    os.makedirs(outdir, exist_ok=True)
    rows = [[i.level, i.row.title, i.row.res, i.row.address, i.text,
             num(i.row.prev), num(i.row.new), None if i.cons is None else round(i.cons, 4),
             os.path.basename(i.row.src), i.row.line] for i in issues]
    if not rows:
        rows = [["", "Замечаний нет", "", "", "", None, None, None, "", None]]
    sheets = [{
        "name": "Проверка",
        "headers": ["Уровень", "Подъезд", "Ресурс", "Адрес", "Что не так", "Было", "Стало", "Расход", "Файл", "Строка"],
        "rows": rows, "widths": [16, 26, 8, 36, 60, 11, 11, 10, 20, 8],
        "fill": lambda v: FILL_ERR if v[0] == "ОШИБКА" else (FILL_WARN if v[0] else None),
    }]
    doubt = [i for i in issues if i.kind == "doubt"]
    if doubt:
        sheets.append({
            "name": "Сомнительные",
            "headers": ["Файл", "Подъезд", "Ресурс", "Адрес", "Исходное значение", "Перенесено (предыдущее)",
                        "Дата предыдущего", "Строка"],
            "rows": [[os.path.basename(i.row.src), i.row.title, i.row.res, i.row.address, str(i.row.new),
                      num(i.row.prev), to_dt(i.row.prev_date), i.row.line] for i in doubt],
            "widths": [20, 26, 8, 36, 18, 22, 16, 8], "numfmt": {7: "dd.mm.yyyy"},
        })
    path = os.path.join(outdir, "проверка.xlsx")
    save_book(path, sheets)
    return path


# ----------------------------------------------------------------------------
# Сводка
# ----------------------------------------------------------------------------
def run_summary(data, outdir):
    os.makedirs(outdir, exist_ok=True)
    ent = OrderedDict()   # (h,c,e) -> {res: [sum, ok, total]}
    for t in data.titles():
        key = parse_title(t)
        acc = ent.setdefault(key, {})
        for res, rows in data.entr[t].items():
            a = acc.setdefault(res, [0.0, 0, 0])
            for r in rows:
                a[2] += 1
                c = cons_of(r)
                if c is not None:
                    a[0] += c
                    a[1] += 1

    def roll(keyfn):
        groups = OrderedDict()
        for key, acc in ent.items():
            g = groups.setdefault(keyfn(key), {})
            for res, (s, ok, tot) in acc.items():
                x = g.setdefault(res, [0.0, 0, 0])
                x[0] += s
                x[1] += ok
                x[2] += tot
        return groups

    def table(groups, labels):
        rows, tot = [], {}
        for key, acc in groups.items():
            vals, miss = [], 0
            for res in RESOURCES:
                if res in acc:
                    s, ok, n = acc[res]
                    vals.append(round(s, 4))
                    miss += n - ok
                    t = tot.setdefault(res, [0.0, 0])
                    t[0] += s
                    t[1] += n - ok
                else:
                    vals.append(None)
            rows.append(list(key) + vals + [miss])
        rows.append(["Итого"] + [""] * (len(labels) - 1) +
                    [round(tot[r][0], 4) if r in tot else None for r in RESOURCES] +
                    [sum(v[1] for v in tot.values())])
        return rows

    numcols = {len(["a"] * 3) + i + 1: "#,##0.000" for i in range(3)}
    sheets = [
        {"name": "Подъезды", "headers": ["Дом", "Корпус", "Подъезд"] + RESOURCES + ["Без показаний"],
         "rows": table(OrderedDict((k, v) for k, v in roll(lambda k: k).items()), ["Дом", "Корпус", "Подъезд"]),
         "widths": [8, 9, 10, 12, 12, 12, 14], "numfmt": numcols, "bold_last": True},
        {"name": "Корпуса", "headers": ["Дом", "Корпус"] + RESOURCES + ["Без показаний"],
         "rows": table(roll(lambda k: (k[0], k[1])), ["Дом", "Корпус"]),
         "widths": [8, 9, 12, 12, 12, 14], "numfmt": {i: "#,##0.000" for i in (3, 4, 5)}, "bold_last": True},
        {"name": "Дома", "headers": ["Дом"] + RESOURCES + ["Без показаний"],
         "rows": table(roll(lambda k: (k[0],)), ["Дом"]),
         "widths": [8, 12, 12, 12, 14], "numfmt": {i: "#,##0.000" for i in (2, 3, 4)}, "bold_last": True},
    ]
    path = os.path.join(outdir, "сводка.xlsx")
    save_book(path, sheets)
    return path


# ----------------------------------------------------------------------------
# Обратная сборка для бухгалтерии
# ----------------------------------------------------------------------------
def write_assembled(path, title_rows):
    """title_rows: [(title, [Row])] — блоки подъездов в формате исходной выгрузки."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист1"
    spec = [("left", None), ("left", None), ("left", None), ("center", DATEF),
            ("right", NUM), ("center", DATEF), ("right", NUM)]
    r = 1
    for title, rows in title_rows:
        ws.cell(r, 1, title).font = F_HEAD
        r += 1
        for row in rows:
            vals = [row.address, row.res, row.serial, to_dt(row.prev_date), row.prev,
                    to_dt(row.new_date), accepted_reading(row)]   # непринятые («?», ошибки) — пусто
            for j, v in enumerate(vals):
                style_cell(ws.cell(r, j + 1, v), *spec[j])
            r += 1
    for k, w in zip("ABCDEFG", (35, 5, 16.84, 9.85, 13.56, 11.28, 14.85)):
        ws.column_dimensions[k].width = w
    wb.save(path)


def run_assemble(data, outdir):
    os.makedirs(outdir, exist_ok=True)
    made = []
    for res in res_order({r.res for r in data.rows()}):
        if res == "ОТ":   # как в исходных выгрузках: отопление — отдельный файл на каждый дом
            houses = OrderedDict()
            for t in data.titles():
                if res in data.entr[t]:
                    houses.setdefault(parse_title(t)[0], []).append((t, data.entr[t][res]))
            for h, blocks in houses.items():
                p = os.path.join(outdir, "ОТ_дом_%s_сводная.xlsx" % h)
                write_assembled(p, blocks)
                made.append(p)
        else:
            blocks = [(t, data.entr[t][res]) for t in data.titles() if res in data.entr[t]]
            p = os.path.join(outdir, "%s_сводная.xlsx" % res)
            write_assembled(p, blocks)
            made.append(p)
    return made


# ----------------------------------------------------------------------------
# Следующий месяц
# ----------------------------------------------------------------------------
def roll_forward(data, new_date):
    out, carried = Data(), 0
    for t in data.titles():
        for res in res_order(data.entr[t]):
            for r in data.entr[t][res]:
                n = accepted_reading(r)
                if n is not None:
                    prev, prev_date = n, r.new_date
                else:       # показание не принято: переносим старое, период растянется
                    prev, prev_date = r.prev, r.prev_date
                    carried += 1
                out.add(Row(title=t, res=res, address=r.address, serial=r.serial,
                            prev_date=prev_date, prev=prev, new_date=new_date, new=None))
    return out, carried


def add_missing(data, sources, new_date):
    """Добавить из свежих выгрузок квартиры, которых нет в таблицах. Серийные номера существующих
    квартир сравниваются с выгрузкой: различие только сообщается (ничего не меняется)."""
    have = {(r.title, r.res, r.address.lower()): r for r in data.rows()}
    added, new_ent, touched, changes = 0, set(), set(), []
    for f in collect_sources(sources):
        for r in parse_export(f):
            key = (r.title, r.res, r.address.lower())
            if key in have:
                old, new = ser(have[key].serial), ser(r.serial)
                if old and new and old.casefold() != new.casefold() and key not in {c[0] for c in changes}:
                    changes.append((key, old, new))
                continue
            if r.title not in data.entr:
                new_ent.add(r.title)
            row = Row(title=r.title, res=r.res, address=r.address, serial=r.serial,
                      prev_date=r.prev_date, prev=r.prev, new_date=new_date, new=None)
            data.add(row)
            have[key] = row
            touched.add((r.title, r.res))
            added += 1
    for t, res in touched:      # новые квартиры встают на своё место по номеру
        data.entr[t][res].sort(key=lambda x: nat(x.address))
    return added, new_ent, changes


def print_serial_changes(changes, limit=40):
    if not changes:
        return
    print("\n" + "-" * 60)
    print("СМЕНА СЕРИЙНОГО №: %d. В новых таблицах оставлен прежний номер — проверьте, не заменён ли счётчик "
          "(при замене предыдущее показание нужно уточнить)." % len(changes))
    for (title, res, address), old, new in changes[:limit]:
        print("  %s | %s | %s: в таблице %s, в выгрузке %s" % (slug(title), res, short(address), old, new))
    if len(changes) > limit:
        print("  ... и ещё %d" % (len(changes) - limit))


def default_next_date(data):
    """20-е число месяца, следующего за месяцем дат съёма в таблицах (самая частая дата съёма)."""
    ds = [to_dt(r.new_date) for r in data.rows() if to_dt(r.new_date)]
    base = Counter(ds).most_common(1)[0][0] if ds else datetime.today()
    y, m = (base.year + 1, 1) if base.month == 12 else (base.year, base.month + 1)
    return datetime(y, m, NEXT_DAY)


def pick_date(a, data):
    if a.date is None:
        a.date = default_next_date(data)
        print("Дата съёма нового месяца: %s (%d-е число следующего месяца)"
              % (a.date.strftime("%d.%m.%Y"), NEXT_DAY))


def default_next_folder(paths, d):
    name = "подъезды_" + d.strftime("%Y-%m")
    if os.path.isdir(paths[0]):          # дали папку -> новая папка рядом
        return os.path.join(os.path.dirname(base_dir(paths)), name)
    return os.path.join(base_dir(paths), name)   # дали файлы -> подпапка рядом с ними


# ----------------------------------------------------------------------------
# Команды
# ----------------------------------------------------------------------------
def cmd_split(a):
    data = Data()
    for f in collect_sources(a.sources):
        for r in parse_export(f):
            data.add(r)
    if not data.entr:
        sys.exit("В файлах не нашлось строк с показаниями.")
    if a.date:
        for r in data.rows():
            r.new_date = a.date
    first = next(iter(data.rows())).new_date
    out = a.out or ("подъезды_" + first.strftime("%Y-%m") if first else "подъезды")
    ensure_out(out, a.overwrite)
    n = write_folder(data, out)
    cnt = sum(1 for _ in data.rows())
    print("Готово: %d файлов (подъездов), %d строк. Папка: %s" % (n, cnt, out))
    return out


def cmd_check(a):
    data, n = load_inputs(a.sources)
    outdir = a.out or default_report_dir(a.sources)
    ensure_report_dir(outdir, a.overwrite, ("check",))
    issues = run_check(data)
    print_issues(issues)
    print("Отчёт: %s" % write_check_report(issues, outdir))


def cmd_summary(a):
    data, _ = load_inputs(a.sources)
    outdir = a.out or default_report_dir(a.sources)
    ensure_report_dir(outdir, a.overwrite, ("summary",))
    print("Сводка: %s" % run_summary(data, outdir))


def cmd_assemble(a):
    data, _ = load_inputs(a.sources)
    outdir = a.out or default_report_dir(a.sources)
    ensure_report_dir(outdir, a.overwrite, ("assemble",))
    for p in run_assemble(data, outdir):
        print("Сводная таблица: %s" % p)


def do_next(a, data):
    new, carried = roll_forward(data, a.date)
    added, new_ent, a.serial_changes = (0, set(), [])
    if a.add:
        added, new_ent, a.serial_changes = add_missing(new, a.add, a.date)
    out = a.out
    ensure_out(out, a.overwrite)
    n = write_folder(new, out)
    cnt = sum(1 for _ in new.rows())
    print("Новый месяц (%s): %d файлов, %d строк. Папка: %s" % (a.date.strftime("%d.%m.%Y"), n, cnt, out))
    if carried:
        print("  Перенесено без нового показания (пусто/ошибка): %d строк — расход посчитается за два периода." % carried)
    if added:
        print("  Добавлено новых квартир из выгрузки: %d" % added)
    for t in sorted(new_ent, key=title_key):
        print("  Новый подъезд: %s" % t)


def cmd_next(a):
    data, _ = load_inputs(a.sources)
    pick_date(a, data)
    a.out = a.out or default_next_folder(a.sources, a.date)
    outdir = default_report_dir(a.sources)
    ensure_out(a.out, a.overwrite)                         # все проверки «не затираем ли чужое» — до записи
    ensure_report_dir(outdir, a.overwrite, ("check",))
    issues = run_check(data)
    report = write_check_report(issues, outdir)
    do_next(a, data)
    print_issues(issues, final=True)
    print_serial_changes(a.serial_changes)
    print("Отчёт: %s" % report)


def cmd_close(a):
    data, _ = load_inputs(a.sources)
    pick_date(a, data)
    a.out = a.out or default_next_folder(a.sources, a.date)
    outdir = default_report_dir(a.sources)
    ensure_out(a.out, a.overwrite)                         # все проверки «не затираем ли чужое» — до записи
    ensure_report_dir(outdir, a.overwrite, ("check", "summary", "assemble"))
    issues = run_check(data)
    report = write_check_report(issues, outdir)
    for p in run_assemble(data, outdir):
        print("Сводная таблица: %s" % p)
    print("Сводка: %s" % run_summary(data, outdir))
    do_next(a, data)
    print_issues(issues, final=True)      # ошибки — в самом конце, чтобы их было видно
    print_serial_changes(a.serial_changes)
    print("Отчёт: %s" % report)


def build_parser():
    ap = argparse.ArgumentParser(description="Показания приборов учёта: подъезды, проверка, сводка, новый месяц.")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("split", help="выгрузки -> файлы по подъездам")
    p.add_argument("sources", nargs="+", help="папка с выгрузками или сами файлы .xls/.xlsx")
    p.add_argument("--out", help="куда сохранить (по умолчанию подъезды_ГГГГ-ММ)")
    p.add_argument("--date", type=date_arg, help="поставить эту дату съёма вместо даты из выгрузки")
    p.add_argument("--overwrite", action="store_true", help="писать в уже существующую папку")
    p.set_defaults(fn=cmd_split)

    for name, fn, hlp in (("check", cmd_check, "проверить заполненные файлы подъездов"),
                          ("summary", cmd_summary, "сводка расхода по подъездам/корпусам/домам"),
                          ("assemble", cmd_assemble, "собрать общие таблицы для бухгалтерии")):
        p = sub.add_parser(name, help=hlp)
        p.add_argument("sources", nargs="+", help="папка с таблицами подъездов или сами файлы .xlsx")
        p.add_argument("--out", help="куда сохранить (по умолчанию подпапка %s рядом с таблицами)" % REPORT_DIR)
        p.add_argument("--overwrite", action="store_true", help="заменить уже существующие отчёты")
        p.set_defaults(fn=fn)

    for name, fn, hlp in (("next", cmd_next, "файлы на следующий месяц из таблиц прошлого"),
                          ("close", cmd_close, "проверка + сводка + сборка + следующий месяц")):
        p = sub.add_parser(name, help=hlp)
        p.add_argument("sources", nargs="+",
                       help="заполненные таблицы прошлого месяца: папка или файлы подъездов (.xlsx)")
        p.add_argument("--date", type=date_arg,
                       help="дата съёма нового месяца, ДД.ММ.ГГГГ (по умолчанию: %d-е число следующего месяца)" % NEXT_DAY)
        p.add_argument("--out", help="куда сохранить файлы нового месяца (по умолчанию подъезды_ГГГГ-ММ рядом)")
        p.add_argument("--add", nargs="+", metavar="ВЫГРУЗКА",
                       help="свежие выгрузки (.xls/.xlsx или папка): добавить квартиры, которых нет в таблицах")
        p.add_argument("--overwrite", action="store_true", help="писать в уже существующую папку")
        p.set_defaults(fn=fn)
    return ap


def main():
    a = build_parser().parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
